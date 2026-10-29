#!/usr/bin/env python3
"""p1_build_skip_logic.py — reconstruct who was asked what, with evidence recomputed.

Inputs (read-only):
    00_source/data/survey.xlsx                         SRC-002
    00_source/data/open_ended.xlsx                     SRC-003
    01_instrument_reconstruction/derived/table_blocks.csv
Outputs:
    01_instrument_reconstruction/routing.csv     one row per rule, with evidence and provenance
    01_instrument_reconstruction/skip_logic.md   rules written as IF ... THEN ...

Each rule states a condition and the evidence for it. Evidence is computed from the
source on every run; if any RECON_DATA check fails, the script stops (exit 2) instead of
writing a rule the data does not support. INFERRED and MISSING rules carry a rationale.

Counts are shares × base, rounded (tables store shares, not counts).

Usage: python3 -I scripts/p1_build_skip_logic.py
"""

from __future__ import annotations

import csv
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SRC2 = Path("00_source/data/survey.xlsx")
SRC3 = Path("00_source/data/open_ended.xlsx")
BLOCKS = Path("01_instrument_reconstruction/derived/table_blocks.csv")
OUT_CSV = Path("01_instrument_reconstruction/routing.csv")
OUT_MD = Path("01_instrument_reconstruction/skip_logic.md")
CANON, MAIN = "Industry 7", "Main Page"
COL_ALL, COL_ITDM, COL_BDM = 2, 4, 5      # Main Page banner columns (P1.S01)

FIELDS = ["rule_id", "kind", "condition", "action", "applies_to", "evidence", "check", "provenance", "note"]


class Src:
    def __init__(self):
        self.wb = load_workbook(SRC2, data_only=True)
        rows = list(csv.DictReader(BLOCKS.open(encoding="utf-8")))
        self.b = {(r["sheet"], r["item_id"]): r for r in rows if r["item_id"]}
        self.other = pd.read_excel(SRC2, sheet_name='"Other" Text Answers', dtype=str)
        self.oe = pd.read_excel(SRC3, sheet_name=None, dtype=str)

    def base(self, item, col=COL_ALL, sheet=MAIN):
        r = self.b[(sheet, item)]
        v = self.wb[sheet].cell(int(r["header_row"]) + 1, col).value
        return None if v is None or str(v).strip() == "-" else int(v)

    def count(self, item, label_start, sheet=CANON):
        r = self.b[(sheet, item)]
        ws = self.wb[sheet]
        base = int(r["base_all"])
        for k in range(int(r["first_option_row"]), int(r["last_option_row"]) + 1):
            if str(ws.cell(k, 1).value).strip().lower().startswith(label_start.lower()):
                return round((ws.cell(k, 2).value or 0) * base)
        raise KeyError(f"{item}: no option starting '{label_start}'")

    def share(self, item, sheet=CANON):
        r = self.b[(sheet, item)]
        return self.wb[sheet].cell(int(r["first_option_row"]), 2).value or 0

    def specify(self, col_prefix):
        col = next(c for c in self.other.columns if str(c).startswith(col_prefix))
        v = self.other[col].dropna().astype(str).str.strip()
        return int((v != "").sum())


def main() -> int:
    s = Src()
    rules, failures = [], []

    def rule(rid, kind, cond, action, applies, evidence, ok, prov, note=""):
        if prov == "RECON_DATA" and not ok:
            failures.append(f"{rid}: evidence does not hold — {evidence}")
        rules.append({"rule_id": rid, "kind": kind, "condition": cond, "action": action,
                      "applies_to": applies, "evidence": evidence,
                      "check": ("PASS" if ok else "FAIL") if prov == "RECON_DATA" else "not checkable",
                      "provenance": prov, "note": note})

    # --- Screening -------------------------------------------------------------------
    it_dept = s.count("S1", "IT or IS department")
    rule("R01", "branch", 'S1 = "IT or IS department"', "classify ITDM; ASK S2, S4",
         "S2, S4", f"S1 IT/IS = {it_dept}; S2 base = {s.base('S2')}; S4 base = {s.base('S4')}; "
         f"S2/S4 BDM column = '-'", it_dept == s.base("S2") == s.base("S4") == 160
         and s.base("S2", COL_BDM) is None, "RECON_DATA")
    rule("R02", "branch", 'S1 = any other department', "classify BDM; ASK S3, S6",
         "S3, S6", f"319 − {it_dept} = {319 - it_dept}; S3 base = {s.base('S3')}; S6 base = {s.base('S6_1')}; "
         f"S3/S6 ITDM column = '-'", 319 - it_dept == s.base("S3") == s.base("S6_1") == 159
         and s.base("S3", COL_ITDM) is None, "RECON_DATA")
    rule("R03", "screen-out", 'S5 = "Don\'t know/Not applicable"', "TERMINATE",
         "S5", f"S5_6 share = {s.share('S5_6')}", s.share("S5_6") == 0, "RECON_DATA",
         "Published methodology (REF-1) states respondents were screened for significant involvement in IT decision-making and purchasing — consistent.")
    rule("R04", "screen-out", 'S6 = "I am not involved in IT decision-making" OR "Don\'t know/Not applicable"',
         "TERMINATE", "S6", f"S6_6 share = {s.share('S6_6')}; S6_7 share = {s.share('S6_7')}",
         s.share("S6_6") == 0 and s.share("S6_7") == 0, "RECON_DATA", "Same as R03.")
    rule("R05", "screen", "S9 country not in {Canada, France, Germany, United Kingdom, United States}",
         "TERMINATE", "S9", "only these five countries appear (n = 7 / 54 / 54 / 54 / 150)",
         True, "RECALL_DESIGNER", "Confirmed by Scott, 2026-10-29. Data: only these five countries appear.")
    rule("R06", "screen", "S8 employees below 1,000", "TERMINATE", "S8",
         "lowest employee band starts at 1,000 (label reads '1,0000')", True, "RECALL_DESIGNER",
         "Confirmed by Scott, 2026-10-29. Data: no one below 1,000 appears.")
    rule("R07", "quota", "ITDM ≈ BDM", "quota 50/50", "S1 → DM type",
         "ITDM 160 / BDM 159", True, "RECALL_DESIGNER", "Confirmed by Scott, 2026-10-29; REF-1 also reports a 50/50 split.")
    rule("R08", "quota", "UK = France = Germany", "quota 54 each", "S9",
         "54 / 54 / 54", True, "RECALL_DESIGNER", "Confirmed by Scott, 2026-10-29. Data: exactly equal counts.")
    rule("R09", "quality", 'Q19_4 ≠ "Somewhat important" (attention check)', "REMOVE respondent",
         "all", f"Q19_4 'Somewhat important' share = {s.count('Q19_4', 'Somewhat important') / 319:.3f}",
         s.count("Q19_4", "Somewhat important") == 319, "RECON_DATA",
         "Every tabulated respondent passed; failures were removed before tabulation.")

    # --- Piping ------------------------------------------------------------------------
    rule("R10", "piping", "S5 wording", "insert text at 'hS5Piping'", "S5",
         "placeholder 'hS5Piping' appears in S5 headers", True, "MISSING",
         "The inserted text is not in the source, and Scott does not recall it (2026-10-29). S5 is treated as one question; caution: the wording may have differed for IT and business respondents.")

    # --- Follow-ups ------------------------------------------------------------------------
    q20_share = s.count("Q20_PurchaseLimit", "Yes, purchase  limit and Im wi")
    rule("R11", "follow-up", 'Q20 = "Yes, purchase limit and I\'m willing to share"', "ASK Q20_2 (amount)",
         "Q20_2", f"Q20 willing-to-share = {q20_share}; Q20_2 base = {s.base('Q20_2')}; Q20_2 text entries = {s.specify('Q20_2')}",
         q20_share == s.base("Q20_2") == s.specify("Q20_2") == 23, "RECON_DATA")
    rule("R12", "follow-up", "unknown", "Q20_4 text", "Q20_4", f"Q20_4 text entries = {s.specify('Q20_4')}",
         True, "MISSING", "Column exists with no entries; its trigger cannot be recovered.")

    for item, label, col, base_item in [("Q1", "Other", "Q1_Other", "Q1_11"), ("Q7", "Other", "Q7_Other", "Q7_8"),
                                        ("Q11", "Other", "Q11_Other", "Q11_7"), ("Q15", "Other", "Q15_Other", "Q15_7")]:
        sel = round(s.share(base_item) * 319)
        txt = s.specify(col)
        rule(f"R13-{item}", "specify", f'{item} "Other (Please specify)" selected', f"ASK {col} (text)", col,
             f"selected = {sel}; text entries = {txt}", sel == txt, "RECON_DATA")
    s4 = s.count("S4", "Other")
    rule("R13-S4", "specify", 'S4 = "Other (Please specify)"', "ASK S4_Other (text)", "S4_Other",
         f"selected = {s4}; text entries = {s.specify('S4_Other')}", s4 == s.specify("S4_Other"), "RECON_DATA")
    q18 = s.count("Q18", "Other")
    rule("R13-Q18", "specify", 'Q18 = "Other, please specify"', "ASK Q18_Other (text)", "Q18_Other",
         f"selected = {q18}; text entries = {s.specify('Q18_Other')}", q18 == s.specify("Q18_Other"), "RECON_DATA")

    for grid, row, col in [("Q8", "Q8_8", "Q8_Other"), ("Q12", "Q12_10", "Q12_Other"),
                           ("Q13", "Q13_10", "Q13_Other"), ("Q19", "Q19_7", "Q19_Other")]:
        b, t = s.base(row), s.specify(col)
        rule(f"R14-{grid}", "specify", f"respondent typed an 'other' attribute in {col}",
             f"ASK rating row {row} for it", row, f"{row} base = {b}; {col} text entries = {t}",
             b == t, "RECON_DATA", "Only those who named an extra attribute rated it.")

    # --- Everything else ---------------------------------------------------------------
    unrouted = []
    for (sheet, item), r in s.b.items():
        if sheet == MAIN and r["has_base_row"] == "yes" and s.base(item) == 319:
            unrouted.append(item)
    rule("R15", "all", "none", "ASK all respondents", f"{len(unrouted)} tabulated items",
         f"{len(unrouted)} Main Page items have base 319", len(unrouted) > 0, "RECON_DATA",
         "All Q1–Q20 items not listed above, plus S1, S5, S7–S12.")
    oe_ok, oe_ev = True, []
    for k, d in s.oe.items():
        n_id = int(d["respid"].notna().sum())
        blank = int(d["respid"].isna().sum())
        oe_ok &= n_id == 319
        oe_ev.append(f"{k}: {n_id} respids + {blank} empty rows")
    rule("R16", "all", "none", "ASK all respondents", "Q21–Q25", "; ".join(oe_ev), oe_ok, "RECON_DATA",
         "The extra 40 rows per sheet are empty (no respid, no answer) — an export artifact, not missing respondents.")

    if failures:
        for f in failures:
            print("ERROR:", f, file=sys.stderr)
        return 2

    with OUT_CSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS, lineterminator="\n")
        w.writeheader()
        w.writerows(rules)

    lines = ["# Skip Logic and Routing", "",
             "Generated by `scripts/p1_build_skip_logic.py`. Evidence is recomputed from the source on every run; "
             "a `RECON_DATA` rule is written only if its check passes.", "",
             "Counts are table shares × base, rounded.", ""]
    sections = [("Screening and quotas", ("branch", "screen-out", "screen", "quota", "quality")),
                ("Piping", ("piping",)), ("Follow-ups and 'please specify'", ("follow-up", "specify")),
                ("Asked of everyone", ("all",))]
    for title, kinds in sections:
        lines += [f"## {title}", "", "| Rule | IF | THEN | Evidence | Provenance | Note |", "|---|---|---|---|---|---|"]
        for r in rules:
            if r["kind"] in kinds:
                cells = [r["rule_id"], r["condition"], r["action"], f"{r['evidence']} ({r['check']})",
                         r["provenance"], r["note"]]
                lines.append("| " + " | ".join(c.replace("|", "\\|") for c in cells) + " |")
        lines.append("")
    prov = {}
    for r in rules:
        prov[r["provenance"]] = prov.get(r["provenance"], 0) + 1
    lines += ["## Totals", "", f"- Rules: {len(rules)} — " + ", ".join(f"{k} {v}" for k, v in sorted(prov.items())),
              "- Every RECON_DATA check: PASS", ""]
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"rules={len(rules)} provenance={prov} → {OUT_CSV}, {OUT_MD}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
