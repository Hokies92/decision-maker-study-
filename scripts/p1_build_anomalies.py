#!/usr/bin/env python3
"""p1_build_anomalies.py — every oddity in the source, with recomputed evidence and a disposition.

Inputs (read-only):
    00_source/data/survey.xlsx, 00_source/data/open_ended.xlsx
    01_instrument_reconstruction/derived/table_blocks.csv
    01_instrument_reconstruction/question_map.csv
Outputs:
    01_instrument_reconstruction/anomalies.csv   one row per anomaly
    01_instrument_reconstruction/anomalies.md    readable list

Nothing in 00_source/ is changed. A disposition says how later phases treat the anomaly
(e.g. "correct label in derived data only"). Each check is recomputed on every run; if a
check that is expected to hold does not, the script stops (exit 2).

Usage: python3 -I scripts/p1_build_anomalies.py
"""

from __future__ import annotations

import csv
import json
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

SRC2 = Path("00_source/data/survey.xlsx")
SRC3 = Path("00_source/data/open_ended.xlsx")
RD = Path("01_instrument_reconstruction")
BLOCKS, QMAP = RD / "derived/table_blocks.csv", RD / "question_map.csv"
OUT_CSV, OUT_MD = RD / "anomalies.csv", RD / "anomalies.md"
CANON, MAIN = "Industry 7", "Main Page"
TOL = 1e-9
FIELDS = ["id", "category", "title", "evidence", "check", "affects_analysis", "disposition", "provenance", "cite"]


def main() -> int:
    wb = load_workbook(SRC2, data_only=True)
    blocks = list(csv.DictReader(BLOCKS.open(encoding="utf-8")))
    B = {(b["sheet"], b["item_id"]): b for b in blocks if b["item_id"]}
    qmap = list(csv.DictReader(QMAP.open(encoding="utf-8")))
    out, fail = [], []

    def col_values(sheet, item, col=2):
        b, ws = B[(sheet, item)], wb[sheet]
        return [ws.cell(r, col).value for r in range(int(b["first_option_row"]), int(b["last_option_row"]) + 1)]

    def labels(sheet, item):
        return json.loads(B[(sheet, item)]["options_json"])

    def add(aid, cat, title, evidence, ok, affects, disposition, prov, cite, expect=True):
        if expect and not ok:
            fail.append(f"{aid}: {evidence}")
        out.append({"id": aid, "category": cat, "title": title, "evidence": evidence,
                    "check": "PASS" if ok else "FAIL", "affects_analysis": affects,
                    "disposition": disposition, "provenance": prov, "cite": cite})

    # A01 — Main Page Q11 condensed
    cond = col_values(MAIN, "Q11_1")
    sep = [col_values(CANON, f"Q11_{k}")[0] for k in range(1, 8)]
    ok = len(cond) == 7 and all(abs(a - b) < TOL for a, b in zip(cond, sep))
    add("A01", "layout", "Main Page shows Q11 as one condensed table",
        f"Main Page 'Q11_1' table has {len(cond)} statements and no base row; its 7 values equal Q11_1–Q11_7 on '{CANON}'",
        ok, "no", f"Same data, different layout. Canonical source for Q11 is '{CANON}'.", "RECON_DATA",
        f"SRC-002 Main Page!A{B[(MAIN, 'Q11_1')]['header_row']}")

    # A02 — Main Page Q15 condensed
    cond = col_values(MAIN, "Q15")
    sep = [col_values(CANON, f"Q15_{k}")[0] for k in range(1, 8)]
    ok = len(cond) == 7 and all(abs(a - b) < TOL for a, b in zip(cond, sep)) and (MAIN, "Q15_8") in B
    add("A02", "layout", "Main Page shows Q15 as one condensed table",
        f"Main Page 'Q15' table has {len(cond)} options; values equal Q15_1–Q15_7 on '{CANON}'; Q15_8 shown separately",
        ok, "no", f"Same data, different layout. Canonical source for Q15 is '{CANON}'.", "RECON_DATA",
        f"SRC-002 Main Page!A{B[(MAIN, 'Q15')]['header_row']}")

    # A03 — Sheet1 duplicates Main Page Q15
    s1 = wb["Sheet1"]
    s1_rows = [(s1.cell(r, 6).value, s1.cell(r, 7).value, s1.cell(r, 9).value, s1.cell(r, 10).value) for r in range(9, 16)]
    mb = B[(MAIN, "Q15")]
    m = wb[MAIN]
    main_rows = [(m.cell(r, 1).value, m.cell(r, 2).value, m.cell(r, 4).value, m.cell(r, 5).value)
                 for r in range(int(mb["first_option_row"]), int(mb["last_option_row"]) + 1)]
    ok = len(s1_rows) == len(main_rows) and all(
        str(a[0]).strip() == str(b[0]).strip() and all(abs(x - y) < TOL for x, y in zip(a[1:], b[1:]))
        for a, b in zip(s1_rows, main_rows))
    add("A03", "duplicate", "'Sheet1' is a copy of the Main Page Q15 table",
        "Sheet1 rows 7–15: Q15 header, base 319/160/159, 7 options; every label and value (All, ITDM, BDM) equals Main Page Q15",
        ok, "no", "Working copy (likely for a chart). Not used as a source; no unique data.", "RECON_DATA",
        "SRC-002 Sheet1!F7")

    # A04 — Q6 bucket label
    q6 = dict(zip(labels(CANON, "Q6"), col_values(CANON, "Q6")))
    zero_to_five = sum(v for k, v in q6.items() if k.isdigit() and int(k) <= 5)
    one_to_five = sum(v for k, v in q6.items() if k.isdigit() and 1 <= int(k) <= 5)
    bucket = dict(zip(labels(CANON, "Q6_bucketed"), col_values(CANON, "Q6_bucketed")))["1-5"]
    ok = abs(bucket - zero_to_five) < 1e-6 and abs(bucket - one_to_five) > 0.01
    add("A04", "label error", "Q6 bucket '1-5' actually holds answers 0–5",
        f"'1-5' share = {bucket:.4f}; sum of answers 0–5 = {zero_to_five:.4f}; sum of 1–5 only = {one_to_five:.4f}; "
        f"answer 0 alone = {q6['0']:.4f}",
        ok, "yes", "Treat the bucket as '0-5' in derived data (Phase 2). Source label left unchanged.", "RECON_DATA",
        f"SRC-002 {CANON}!A{B[(CANON, 'Q6_bucketed')]['header_row']}")

    # A05 — Q6 outlier
    n80 = round(q6.get("80", 0) * 319)
    add("A05", "outlier", "Q6: one respondent reports 80 appointments a week",
        f"answer '80' share = {q6.get('80', 0):.4f} (= {n80} respondent); next highest answer is 20",
        n80 == 1, "yes", "Keep as recorded; flag as outlier. Use medians or buckets, not means, for Q6.",
        "RECON_DATA", f"SRC-002 {CANON}!A{B[(CANON, 'Q6')]['header_row']}")

    # A06 — Q20_4 empty
    other = pd.read_excel(SRC2, sheet_name='"Other" Text Answers', dtype=str)
    col = next(c for c in other.columns if str(c).startswith("Q20_4"))
    n = int(other[col].dropna().astype(str).str.strip().ne("").sum())
    add("A06", "empty field", "Text column Q20_4 has no entries",
        f"Q20_4 non-blank entries = {n}; trigger not recoverable (skip_logic R12)", n == 0, "no",
        "Kept in the map as MISSING purpose; no data to carry.", "MISSING", "SRC-002 \"Other\" Text Answers!N1")

    # A07 — S5 piping
    piped = sum(1 for q in qmap if "hS5Piping" in q["piping_token"])
    add("A07", "piping", "S5 wording contains an unresolved placeholder",
        f"'hS5Piping' appears in {piped} S5 items; inserted text not in source; Scott does not recall (2026-10-29)",
        piped == 6, "yes", "S5 treated as one question; caution: wording may have differed for IT and business respondents.",
        "MISSING", f"SRC-002 {CANON}!A{B[(CANON, 'S5_1')]['header_row']}")

    # A08 — employee label typo
    lab = labels(CANON, "S8_Bucketed")[0]
    add("A08", "label error", "Employee band reads '1,0000 to 4,999 employees'",
        f"first S8 band label = '{lab}'", lab.startswith("1,0000"), "no",
        "Read as 1,000–4,999 (confirmed screen: 1,000+ employees, Scott 2026-10-29). Correct in derived data only.",
        "RECON_DATA", f"SRC-002 {CANON}!A{int(B[(CANON, 'S8_Bucketed')]['first_option_row'])}")

    # A09 — stems without closing quote
    noclose = [q["item_id"] for q in qmap if q["stem_complete"] == "no"]
    add("A09", "wording", "Question wording without a closing quote (may be truncated)",
        f"items: {', '.join(noclose)}", True, "no",
        "Wording kept VERBATIM as it appears; sibling items carry the full stem.", "RECON_DATA",
        "question_map.csv", expect=False)

    # A10 — stems truncated with an ellipsis in source
    ell = sorted({q["thread_id"] for q in qmap if re.search(r"(\.\.\.|…)\s*$", q["stem"])})
    add("A10", "wording", "Question wording cut off with '...' in the source",
        f"threads: {', '.join(t[2] + str(int(t[3:])) for t in ell) or 'none'}", True, "no",
        "Full wording not recoverable from source; kept VERBATIM with the ellipsis.", "MISSING",
        "question_map.csv", expect=False)

    # A11 — open-ended empty rows
    oe = pd.read_excel(SRC3, sheet_name=None, dtype=str)
    empties = {k: int(d["respid"].isna().sum()) for k, d in oe.items()}
    add("A11", "export artifact", "Open-ended sheets have 40 empty rows each",
        "; ".join(f"{k}: {v}" for k, v in empties.items()) + " (no respid, no answer)",
        all(v == 40 for v in empties.values()), "no",
        "Ignore empty rows; 319 respondents per sheet (skip_logic R16).", "RECON_DATA", "SRC-003 all sheets")

    # A12 — 'Grouping' column empty
    grp = {k: int(d["Grouping"].dropna().astype(str).str.strip().ne("").sum()) for k, d in oe.items()}
    add("A12", "empty field", "Open-ended 'Grouping' column is empty",
        "; ".join(f"{k}: {v} entries" for k, v in grp.items()), all(v == 0 for v in grp.values()), "no",
        "No grouping carried; Phase 5 codes answers from scratch.", "RECON_DATA", "SRC-003 column B")

    # A13 — label typos in options (double spaces, missing apostrophes)
    typo = sorted({f"{b['item_id']}: '{o}'" for b in blocks if b["sheet"] == CANON
                   for o in json.loads(b["options_json"]) if "  " in o or re.search(r"\b(Im|Dont)\b", o)})
    add("A13", "label typo", "Option labels with double spaces or missing apostrophes",
        f"{len(typo)} labels, e.g. {typo[0] if typo else '—'}", len(typo) > 0, "no",
        "Kept VERBATIM in reconstruction; normalized display text allowed in derived data.", "RECON_DATA",
        f"SRC-002 {CANON}", expect=False)

    # A14 — empty 'Other' industry category
    ind = dict(zip(labels(CANON, "S12_Industry7"), col_values(CANON, "S12_Industry7")))
    oth = next(v for k, v in ind.items() if k.lower().startswith("other"))
    add("A14", "empty category", "Industry-7 recode has an 'Other or don't know' category with nobody in it",
        f"'Other or Dont know/not applicable' share = {oth}", oth == 0, "no",
        "Category kept; empty. Consistent with the 39-category list (its 'Don't know' share is also near zero).",
        "RECON_DATA", f"SRC-002 {CANON}!A{B[(CANON, 'S12_Industry7')]['header_row']}")

    # A15 — S6 labels carry a population prefix
    s6 = [q["item_label"] for q in qmap if q["thread_id"] == "T-S06" and q["question_type"] == "multi_select_option"]
    pref = sum(1 for l in s6 if l.startswith("Business Decision Maker - "))
    add("A15", "label quirk", "S6 choice labels repeat 'Business Decision Maker - '",
        f"{pref} of {len(s6)} S6 choice labels start with the population name", pref == len(s6), "no",
        "Kept VERBATIM; display text may drop the prefix in derived data.", "RECON_DATA", "question_map.csv")

    # A16 — significance-test comments
    comment_counts, templates = {}, set()
    for name in ("Main Page", "Industry 7", "Company Size & Revenue", "Sales Engagement", "Sheet1"):
        ws = wb[name]
        cm = [c.comment.text for row in ws.iter_rows() for c in row if c.comment]
        comment_counts[name] = len(cm)
        for t in cm:
            body = re.sub(r"^[^\n]*:\n", "", t)
            templates.add(re.sub(r"\d+(\.\d+)?", "#", body).strip()[:90])
    add("A16", "embedded metadata", "Significance-test results are stored as cell comments",
        f"comments per sheet: {comment_counts}; {len(templates)} wording templates "
        "(significant / not significant / no testing — table too small)",
        sum(comment_counts.values()) > 0, "yes",
        "Carry into Phase 2 as a significance flag per result; not used in Phase 1.", "RECON_DATA",
        "SRC-002 cell comments", expect=False)

    if fail:
        for f in fail:
            print("ERROR:", f, file=sys.stderr)
        return 2

    with OUT_CSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS, lineterminator="\n")
        w.writeheader(); w.writerows(out)

    L = ["# Anomalies", "", "Generated by `scripts/p1_build_anomalies.py`. Evidence recomputed on every run. "
         "Nothing in `00_source/` is changed; dispositions say how later phases handle each item.", "",
         f"**{len(out)} anomalies** — {sum(1 for a in out if a['affects_analysis'] == 'yes')} affect analysis.", "",
         "| ID | What | Evidence | Affects analysis | Disposition | Provenance |", "|---|---|---|---|---|---|"]
    for a in out:
        cells = [a["id"], a["title"], a["evidence"], a["affects_analysis"], a["disposition"], a["provenance"]]
        L.append("| " + " | ".join(str(c).replace("|", "\\|") for c in cells) + " |")
    L.append("")
    OUT_MD.write_text("\n".join(L), encoding="utf-8")
    print(f"anomalies={len(out)} affecting_analysis={sum(1 for a in out if a['affects_analysis'] == 'yes')}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
