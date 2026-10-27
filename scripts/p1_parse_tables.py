#!/usr/bin/env python3
"""p1_parse_tables.py — locate every question block in the tabulated survey file.

Input : 00_source/data/survey.xlsx (SRC-002), read-only.
Output: 01_instrument_reconstruction/derived/table_blocks.csv

A block is defined by the layout observed in SRC-002 (RECON_DATA):
    row r     : header text in column A; column B (first value column) empty
    row r + 1 : usually "Sample Size" with base counts — optional (e.g., the Main Page
                Q11 table has no base row); flagged in `has_base_row`
    rows …    : one answer option per row (column A text, column B populated),
                until a row with an empty column A or column B

Only structure is extracted (where each block sits, its text, its options, its total
base). Result values are not read here; they belong to Phase 2.

Usage: python3 -I scripts/p1_parse_tables.py [--src PATH] [--out PATH]
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

TABLE_SHEETS = [
    "Main Page",
    "Industry 7",
    "Company Size & Revenue",
    "Sales Engagement",
    "Decision Making Stages",
]
BASE_LABEL = "Sample Size"
ITEM_ID = re.compile(r"^\s*((?:Q|S)\d+(?:_[A-Za-z0-9]+)?)")
# e.g. "S2 and S3 combined: ..." → S2_S3_combined (a merged table, not item S2)
COMBINED_ID = re.compile(r"^\s*((?:Q|S)\d+) and ((?:Q|S)\d+) combined", re.I)
FIELDS = [
    "sheet", "block_seq", "header_row", "item_id", "base_question",
    "header_text", "has_base_row", "base_all", "n_options", "first_option_row",
    "last_option_row", "options_json",
]


def cell_text(ws, r: int, c: int) -> str:
    v = ws.cell(r, c).value
    return "" if v is None else str(v).strip()


def is_header(ws, r: int) -> bool:
    return bool(cell_text(ws, r, 1)) and not cell_text(ws, r, 2) \
        and bool(cell_text(ws, r + 1, 1)) and bool(cell_text(ws, r + 1, 2))


def parse_sheet(ws) -> list[dict]:
    blocks: list[dict] = []
    max_row = ws.max_row
    r = 1
    while r < max_row:
        header = cell_text(ws, r, 1)
        if is_header(ws, r):
            has_base = cell_text(ws, r + 1, 1) == BASE_LABEL
            first = r + 2 if has_base else r + 1
            options: list[str] = []
            rr = first
            while rr <= max_row:
                label = cell_text(ws, rr, 1)
                if not label or not cell_text(ws, rr, 2):
                    break
                options.append(label)
                rr += 1
            mc = COMBINED_ID.match(header)
            m = ITEM_ID.match(header)
            item_id = f"{mc.group(1)}_{mc.group(2)}_combined" if mc else (m.group(1) if m else "")
            base_q = re.match(r"(Q|S)\d+", item_id).group(0) if item_id else ""
            base_all = ws.cell(r + 1, 2).value if has_base else None
            blocks.append({
                "sheet": ws.title,
                "block_seq": len(blocks) + 1,
                "header_row": r,
                "item_id": item_id,
                "base_question": base_q,
                "header_text": header,
                "has_base_row": "yes" if has_base else "no",
                "base_all": "" if base_all is None else str(base_all).strip(),
                "n_options": len(options),
                "first_option_row": first if options else "",
                "last_option_row": first + len(options) - 1 if options else "",
                "options_json": json.dumps(options, ensure_ascii=False),
            })
            r = rr
        else:
            r += 1
    return blocks


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--src", type=Path, default=Path("00_source/data/survey.xlsx"))
    ap.add_argument("--out", type=Path,
                    default=Path("01_instrument_reconstruction/derived/table_blocks.csv"))
    args = ap.parse_args(argv)

    wb = load_workbook(args.src, read_only=True, data_only=True)
    missing = [s for s in TABLE_SHEETS if s not in wb.sheetnames]
    if missing:
        print(f"ERROR: expected sheets not found: {missing}", file=sys.stderr)
        return 2
    # read_only worksheets do not support random access well; reload fully.
    wb = load_workbook(args.src, data_only=True)

    rows: list[dict] = []
    for name in TABLE_SHEETS:
        rows += parse_sheet(wb[name])

    args.out.parent.mkdir(parents=True, exist_ok=True)
    with args.out.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS, lineterminator="\n")
        w.writeheader()
        w.writerows(rows)

    by_sheet: dict[str, int] = {}
    for row in rows:
        by_sheet[row["sheet"]] = by_sheet.get(row["sheet"], 0) + 1
    for name in TABLE_SHEETS:
        print(f"{name}: {by_sheet.get(name, 0)} blocks")
    print(f"TOTAL: {len(rows)} blocks → {args.out}")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
