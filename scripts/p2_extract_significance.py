#!/usr/bin/env python3
"""p2_extract_significance.py — turn the significance-test cell comments into flags.

Each comment sits on a question's header row, in the first column of a breakout banner
(column A stands for the All-respondents banner). It records whether that breakout
dimension and that question are related, in one of three wordings:
    "There is a significant relationship between the Column Variable(s) and this Row Variable" → significant
    "There is no significant relationship …"                                                 → not_significant
    "No testing done - table smaller than RxC"                                                → not_tested
The test statistic and threshold are not stated in the source; the report header gives
"Confidence Level = 95%" (Main Page!A10).

Inputs (read-only): SRC-002; table_blocks.csv; cohorts.csv; question_map.csv
Output: 02_data_inventory/significance.csv
Exit 2 if any comment cannot be tied to a block and a dimension, or its wording is unknown.
Sheet1 (duplicate of Main Page Q15, anomaly A03) is not extracted.

Usage: python3 -I scripts/p2_extract_significance.py
"""

from __future__ import annotations

import csv
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

SRC2 = Path("00_source/data/survey.xlsx")
RD = Path("01_instrument_reconstruction")
OUT = Path("02_data_inventory/significance.csv")
SHEETS = ["Main Page", "Industry 7", "Company Size & Revenue", "Sales Engagement", "Decision Making Stages"]
EXPECTED = {"Main Page": 742, "Industry 7": 236, "Company Size & Revenue": 354, "Sales Engagement": 354,
            "Decision Making Stages": 0}
FIELDS = ["sheet", "cell", "block_item_id", "thread_id", "dimension_id", "state", "table_size", "comment"]


def main() -> int:
    wb = load_workbook(SRC2, data_only=True)
    blocks = {(b["sheet"], int(b["header_row"])): b for b in
              csv.DictReader((RD / "derived/table_blocks.csv").open(encoding="utf-8"))}
    cohorts = list(csv.DictReader((RD / "cohorts.csv").open(encoding="utf-8")))
    thread_of = {q["item_id"]: q["thread_id"] for q in csv.DictReader((RD / "question_map.csv").open(encoding="utf-8"))}
    thread_of.update({"Q15": "T-Q15", "DM_TYPE": "T-S01"})
    col_dim = {}
    for c in cohorts:
        if c["status"] == "tabulated" and c["column"]:
            col_dim.setdefault((c["sheet"], column_index_from_string(c["column"])), c["dimension_id"])
    rows, errors, counts = [], [], Counter()
    for sheet in SHEETS:
        ws = wb[sheet]
        for row in ws.iter_rows():
            for cell in row:
                if not cell.comment:
                    continue
                counts[sheet] += 1
                text = re.sub(r"^[^\n]*:\n", "", cell.comment.text).strip()
                body = re.sub(r"^Column %:\s*", "", text)
                if body.startswith("There is a significant relationship"):
                    state, size = "significant", ""
                elif body.startswith("There is no significant relationship"):
                    state, size = "not_significant", ""
                elif m := re.match(r"No testing done - table smaller than (\d+x\d+)", body):
                    state, size = "not_tested", m.group(1)
                else:
                    errors.append(f"{sheet}!{cell.coordinate}: unknown wording {body[:60]!r}"); continue
                b = blocks.get((sheet, cell.row))
                if not b:
                    errors.append(f"{sheet}!{cell.coordinate}: comment not on a block header row"); continue
                block_id = b["item_id"] or "DM_TYPE"
                dim = "ALL" if cell.column == 1 else col_dim.get((sheet, cell.column))
                if not dim:
                    errors.append(f"{sheet}!{cell.coordinate}: column matches no banner"); continue
                rows.append({"sheet": sheet, "cell": cell.coordinate, "block_item_id": block_id,
                             "thread_id": thread_of.get(block_id, ""), "dimension_id": dim, "state": state,
                             "table_size": size, "comment": body})
    for s, n in EXPECTED.items():
        if counts[s] != n:
            errors.append(f"{s}: {counts[s]} comments, expected {n}")
    if any(not r["thread_id"] for r in rows):
        errors.append("some flags have no thread")
    if errors:
        for e in errors[:20]:
            print("ERROR:", e, file=sys.stderr)
        return 2
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS, lineterminator="\n")
        w.writeheader(); w.writerows(rows)
    st = Counter(r["state"] for r in rows)
    sig = Counter(r["dimension_id"] for r in rows if r["state"] == "significant")
    print(f"flags={len(rows)} states={dict(st)}")
    print("significant by dimension:", dict(sig.most_common()))
    return 0


if __name__ == "__main__":
    sys.exit(main())
