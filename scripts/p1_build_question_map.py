#!/usr/bin/env python3
"""p1_build_question_map.py — build the question map: one thread per base question.

Inputs (read-only):
    00_source/data/survey.xlsx                         SRC-002
    00_source/data/open_ended.xlsx                     SRC-003
    01_instrument_reconstruction/derived/table_blocks.csv   (P1.S01)
    01_instrument_reconstruction/thread_annotations.csv     (authored, INFERRED)
Outputs:
    01_instrument_reconstruction/question_map.csv   one row per item
    01_instrument_reconstruction/question_map.md    one row per thread + item counts

Canonical sheet for tabulated items is "Industry 7": it carries every tabulated item
in uncondensed form (P1.S01). Main Page layout is recorded per item, not used as source.

Provenance per field:
    stem, item_label, header_text       VERBATIM (cited by file!sheet!cell)
    question_type                       RECON_DATA (evidence column states the rule)
    derived_from                        RECON_DATA or INFERRED (evidence column)
    construct, section                  INFERRED (from thread_annotations.csv)

Usage: python3 -I scripts/p1_build_question_map.py
"""

from __future__ import annotations

import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

SRC2 = Path("00_source/data/survey.xlsx")
SRC3 = Path("00_source/data/open_ended.xlsx")
BLOCKS = Path("01_instrument_reconstruction/derived/table_blocks.csv")
ANNOT = Path("01_instrument_reconstruction/thread_annotations.csv")
OUT_CSV = Path("01_instrument_reconstruction/question_map.csv")
OUT_MD = Path("01_instrument_reconstruction/question_map.md")

CANON = "Industry 7"
OTHER_SHEET = '"Other" Text Answers'
SUM_TOL = 0.002

ID_PREFIX = re.compile(r"^\s*((?:Q|S)\d+(?:_[A-Za-z0-9]+)?)\s*:?\s*")
COMBINED = re.compile(r"^\s*(?:Q|S)\d+ and (?:Q|S)\d+ combined\s*:\s*", re.I)
PIPE = re.compile(r"'?(h[SQ]\d+Piping\w*)'?")

# Derived items: (item_id) -> (derived_from, provenance, evidence)
DERIVED = {
    "DM_TYPE": ("S1", "INFERRED",
                "No source ID. ITDM share (0.5016, n=160) equals S1 'IT or IS department' share; "
                "BDM = all other S1 departments. Used as the ITDM/BDM stakeholder split."),
    "S2_S3_combined": ("S2;S3", "RECON_DATA",
                       "Header says 'S2 and S3 combined'; same four tiers; base 319 = 160 + 159."),
    "S9_Region": ("S9_Country", "RECON_DATA",
                  "Two regions; North America 157 = US 150 + Canada 7; Europe 162 = UK + France + Germany (54 each)."),
    "S12_Industry7": ("S12", "RECON_DATA", "Item name 'Industry7'; 39 industries collapsed to 7."),
    "Q6_bucketed": ("Q6", "RECON_DATA", "Item name 'bucketed'; 14 counts collapsed to 5 ranges."),
    "S8_Bucketed": ("S8 (raw, not tabulated)", "INFERRED",
                    "Item name 'Bucketed'; the raw S8 answer is not in any source file."),
}
ATTENTION_CHECK = re.compile(r"for quality control,\s*please select", re.I)

FIELDS = [
    "thread_id", "item_id", "source_label_id", "item_seq", "item_role", "question_type",
    "type_evidence", "type_prov", "stem", "stem_complete", "item_label", "text_prov",
    "piping_token", "header_text", "source_cite", "n_options", "options_ref", "base_all",
    "derived_from", "derived_prov", "derived_evidence", "main_page_layout",
]


def thread_of(base_q: str) -> str:
    return f"T-{base_q[0]}{int(base_q[1:]):02d}"


def split_header(header: str) -> tuple[str, str, str, bool]:
    """Return (source_label_id, label, stem, stem_complete). All VERBATIM substrings."""
    m = COMBINED.match(header) or ID_PREFIX.match(header)
    sid = ""
    rest = header
    if m:
        sid = header[: m.end()].strip().rstrip(":").strip()
        rest = header[m.end():]
    if rest.count('"') == 1 and rest.rstrip().endswith('"'):
        rest = rest.rstrip()[:-1]          # stray closing quote with no opening quote
    if rest.count('"') >= 2:
        before, _, after = rest.partition('"')
        stem = after[:-1] if after.endswith('"') else after
        complete = after.endswith('"')
        label = before.strip().rstrip("-").strip()
    elif " - " in rest:
        label, _, stem = rest.partition(" - ")
        complete = False
        label = label.strip()
    else:
        label, stem, complete = "", rest, False
    return sid, label, stem.strip(), complete


def stem_flag(item_id: str, stem: str, complete: bool) -> str:
    if item_id == "DM_TYPE":
        return "n/a"              # table title, not question wording
    return "yes" if complete else "no"


def block_sum(ws, r: dict) -> float:
    if not r["first_option_row"]:
        return 0.0
    vals = [ws.cell(i, 2).value for i in range(int(r["first_option_row"]), int(r["last_option_row"]) + 1)]
    return sum(v for v in vals if isinstance(v, (int, float)))


def main() -> int:
    wb2 = load_workbook(SRC2, data_only=True)
    ws_canon = wb2[CANON]
    blocks = list(csv.DictReader(BLOCKS.open(encoding="utf-8")))
    canon = [b for b in blocks if b["sheet"] == CANON]
    main_by_id = {b["item_id"]: b for b in blocks if b["sheet"] == "Main Page" and b["item_id"]}
    annot = {a["thread_id"]: a for a in csv.DictReader(ANNOT.open(encoding="utf-8"))}

    siblings = defaultdict(list)
    for b in canon:
        if b["base_question"]:
            siblings[b["base_question"]].append(b)

    items: list[dict] = []
    seq = defaultdict(int)

    # 1. Tabulated items (SRC-002, canonical sheet)
    for b in canon:
        header = b["header_text"]
        item_id = b["item_id"] or "DM_TYPE"
        base_q = b["base_question"] or "S1"
        sid, label, stem, complete = split_header(header)
        if item_id == "DM_TYPE":
            label, stem = header, ""
        n = int(b["n_options"])
        s = block_sum(ws_canon, b)
        fam = siblings.get(base_q, [])
        derived = DERIVED.get(item_id)
        if ATTENTION_CHECK.search(header):
            role = "attention_check"
        elif derived:
            role = "derived"
        else:
            role = "asked"

        if n == 1:
            qtype, ev = "multi_select_option", f"single-row table (share selecting = {s:.3f}); one of {sum(1 for x in fam if x['n_options'] == '1')} sibling options"
        elif abs(s - 1) <= SUM_TOL and sum(1 for x in fam if int(x["n_options"]) == n and x["options_json"] == b["options_json"]) >= 2:
            qtype, ev = "grid_row", f"{n} options sum to {s:.3f}; options shared with sibling rows"
        elif abs(s - 1) <= SUM_TOL:
            qtype, ev = "single_select", f"{n} options sum to {s:.3f}"
        else:
            qtype, ev = "unresolved", f"{n} options sum to {s:.3f}"

        if item_id in main_by_id:
            layout = "same" if main_by_id[item_id]["n_options"] == b["n_options"] else \
                f"different ({main_by_id[item_id]['n_options']} options on Main Page)"
        else:
            layout = "absent from Main Page"

        seq[base_q] += 1
        items.append({
            "thread_id": thread_of(base_q), "item_id": item_id, "source_label_id": sid,
            "item_seq": seq[base_q], "item_role": role, "question_type": qtype,
            "type_evidence": ev, "type_prov": "RECON_DATA",
            "stem": stem, "stem_complete": stem_flag(item_id, stem, complete), "item_label": label,
            "text_prov": "VERBATIM", "piping_token": ";".join(PIPE.findall(header)),
            "header_text": header, "source_cite": f"SRC-002 {CANON}!A{b['header_row']}",
            "n_options": n, "options_ref": f"table_blocks.csv: {CANON} row {b['header_row']}",
            "base_all": b["base_all"],
            "derived_from": derived[0] if derived else "", "derived_prov": derived[1] if derived else "",
            "derived_evidence": derived[2] if derived else "", "main_page_layout": layout,
        })

    # 2. "Other, please specify" and follow-up text items (SRC-002 "Other" sheet)
    ws_o = wb2[OTHER_SHEET]
    for c in ws_o[1]:
        header = "" if c.value is None else str(c.value).strip()
        if not header or header == "respid":
            continue
        sid, label, stem, complete = split_header(header)
        base_q = re.match(r"(Q|S)\d+", sid).group(0)
        item_id = sid if sid.endswith("_Other") else f"{sid}_text"
        seq[base_q] += 1
        items.append({
            "thread_id": thread_of(base_q), "item_id": item_id, "source_label_id": sid,
            "item_seq": seq[base_q], "item_role": "asked", "question_type": "open_text_specify",
            "type_evidence": "free-text column keyed by respid", "type_prov": "RECON_DATA",
            "stem": stem, "stem_complete": "yes" if complete else "no", "item_label": label,
            "text_prov": "VERBATIM", "piping_token": ";".join(PIPE.findall(header)),
            "header_text": header, "source_cite": f"SRC-002 {OTHER_SHEET}!{c.coordinate}",
            "n_options": "", "options_ref": "", "base_all": "",
            "derived_from": "", "derived_prov": "", "derived_evidence": "", "main_page_layout": "n/a",
        })

    # 3. Open-ended questions (SRC-003)
    wb3 = load_workbook(SRC3, data_only=True)
    for ws in wb3.worksheets:
        cell = ws["C1"]
        header = str(cell.value).strip()
        sid, label, stem, complete = split_header(header)
        base_q = re.match(r"(Q|S)\d+", sid).group(0)
        seq[base_q] += 1
        items.append({
            "thread_id": thread_of(base_q), "item_id": sid, "source_label_id": sid,
            "item_seq": seq[base_q], "item_role": "asked", "question_type": "open_text",
            "type_evidence": "free-text column keyed by respid, one sheet per question", "type_prov": "RECON_DATA",
            "stem": stem, "stem_complete": "yes" if complete else "no", "item_label": label,
            "text_prov": "VERBATIM", "piping_token": ";".join(PIPE.findall(header)),
            "header_text": header, "source_cite": f"SRC-003 {ws.title}!C1",
            "n_options": "", "options_ref": "", "base_all": "",
            "derived_from": "", "derived_prov": "", "derived_evidence": "", "main_page_layout": "n/a",
        })

    def order(it):
        t = it["thread_id"]
        return (t[2] != "S", int(t[3:]), int(it["item_seq"]))

    items.sort(key=order)
    ids = [i["item_id"] for i in items]
    dupes = sorted({x for x in ids if ids.count(x) > 1})
    if dupes:
        print(f"ERROR: duplicate item_id: {dupes}", file=sys.stderr)
        return 2
    missing_annot = sorted({i["thread_id"] for i in items} - set(annot))
    if missing_annot:
        print(f"ERROR: threads without annotation: {missing_annot}", file=sys.stderr)
        return 2

    with OUT_CSV.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=FIELDS, lineterminator="\n")
        w.writeheader()
        w.writerows(items)

    # Markdown summary — one row per thread
    by_thread = defaultdict(list)
    for it in items:
        by_thread[it["thread_id"]].append(it)
    lines = [
        "# Question Map — thread summary",
        "",
        "Generated by `scripts/p1_build_question_map.py`. Do not edit by hand.",
        "Item-level detail, wording, and citations: `question_map.csv`.",
        "",
        "Provenance: wording `VERBATIM`; question types `RECON_DATA`; construct and section `INFERRED` "
        "(authored in `thread_annotations.csv`, pending Scott's review in P1.S09).",
        "",
        "| Thread | Construct (INFERRED) | Section (INFERRED) | Items | Asked | Derived | Attention check | Types | Stem (VERBATIM, first item) |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    totals = defaultdict(int)
    for t in sorted(by_thread, key=lambda t: (t[2] != "S", int(t[3:]))):
        its = by_thread[t]
        roles = defaultdict(int)
        for i in its:
            roles[i["item_role"]] += 1
            totals[i["item_role"]] += 1
            totals[i["question_type"]] += 0
        types = sorted({i["question_type"] for i in its})
        first = next((i for i in its if i["item_role"] == "asked"), its[0])
        stem = first["stem"].replace("|", "\\|")
        stem = stem if len(stem) <= 90 else stem[:87] + "…"
        a = annot[t]
        lines.append(f"| {t} | {a['construct']} | {a['section']} | {len(its)} | {roles['asked']} | "
                     f"{roles['derived']} | {roles['attention_check']} | {', '.join(types)} | {stem} |")
    type_counts = defaultdict(int)
    for i in items:
        type_counts[i["question_type"]] += 1
    lines += [
        "",
        "## Totals",
        "",
        f"- Threads: {len(by_thread)}",
        f"- Items: {len(items)} (asked {totals['asked']}, derived {totals['derived']}, attention check {totals['attention_check']})",
        "- By type: " + ", ".join(f"{k} {v}" for k, v in sorted(type_counts.items())),
        f"- Items whose stem lacks a closing quote (may be truncated in source): "
        f"{sum(1 for i in items if i['stem_complete'] == 'no')}",
        f"- Items with a piping token: {sum(1 for i in items if i['piping_token'])}",
        "",
    ]
    OUT_MD.write_text("\n".join(lines), encoding="utf-8")
    print(f"threads={len(by_thread)} items={len(items)} → {OUT_CSV}, {OUT_MD}")
    print("types:", dict(sorted(type_counts.items())))
    print("roles:", {k: totals[k] for k in ("asked", "derived", "attention_check")})
    return 0


if __name__ == "__main__":
    sys.exit(main())
