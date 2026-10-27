#!/usr/bin/env python3
"""p0_scan_identifiers.py — screen source data files for respondent-identifying content.

GOVERNANCE D3: no respondent-identifying data is committed in clear text. This scan
runs on candidate files BEFORE they are committed (P0.S04). It never prints cell
values: the report names files, sheets, columns, check categories, and counts only,
so the report itself is safe to commit.

Usage:
    python3 -I scripts/p0_scan_identifiers.py FILE [FILE ...] [--out REPORT.md]

Exit codes: 0 = clear, 3 = findings (stop condition S8; Scott decides case by case),
            2 = usage or read error.

Severity:
    HIGH   — direct identifiers: email, phone, IP, URL/profile link, person-name or
             address headers.
    REVIEW — may identify in context: organization headers, ID-like columns, free-text
             columns (names can appear in prose and cannot be detected reliably),
             file metadata, hidden sheets.
"""

from __future__ import annotations

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

HEADER_RULES: list[tuple[str, str, re.Pattern]] = [
    ("HIGH", "header: person name", re.compile(r"\b(first|last|full|given|sur|respondent)[ _-]?name\b|^name$", re.I)),
    ("HIGH", "header: email", re.compile(r"e-?mail", re.I)),
    ("HIGH", "header: phone", re.compile(r"\b(phone|mobile|cell|tel(ephone)?)\b", re.I)),
    ("HIGH", "header: address", re.compile(r"\b(address|street|zip|postal|postcode)\b", re.I)),
    ("HIGH", "header: IP / device", re.compile(r"\b(ip|ip[ _-]?address|device[ _-]?id)\b", re.I)),
    ("HIGH", "header: profile link", re.compile(r"linkedin|profile[ _-]?url", re.I)),
    ("REVIEW", "header: organization", re.compile(r"\b(company|organi[sz]ation|employer|firm|account)[ _-]?(name)?\b", re.I)),
    ("REVIEW", "header: ID-like", re.compile(r"\b(id|uid|respid|resp[ _-]?id|record[ _-]?id|panel[ _-]?id)\b", re.I)),
]

VALUE_RULES: list[tuple[str, str, re.Pattern]] = [
    ("HIGH", "value: email", re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")),
    ("HIGH", "value: phone", re.compile(r"(?<!\d)(?:\+?\d[\d\s().-]{8,}\d)(?!\d)")),
    ("HIGH", "value: IPv4", re.compile(r"\b(?:\d{1,3}\.){3}\d{1,3}\b")),
    ("HIGH", "value: URL", re.compile(r"https?://|www\.", re.I)),
]

# Decimals, or integers too short to be a phone number (<10 digits).
PURE_NUMBER = r"-?\d*\.\d+(?:[eE]-?\d+)?|-?\d{1,9}"

FREE_TEXT_MIN_MEAN_LEN = 25   # mean characters per non-empty cell
FREE_TEXT_MIN_UNIQUE = 0.5    # share of non-empty cells that are distinct


@dataclass
class Finding:
    file: str
    sheet: str
    column: str
    severity: str
    check: str
    count: int | None  # cells matched; None for header/metadata checks


def read_sheets(path: Path) -> dict[str, pd.DataFrame]:
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return pd.read_excel(path, sheet_name=None, dtype=str)
    if suffix == ".csv":
        return {"(csv)": pd.read_csv(path, dtype=str)}
    raise ValueError(f"unsupported file type: {suffix}")


def workbook_checks(path: Path) -> list[Finding]:
    """Metadata and hidden-structure checks (xlsx/xlsm only). Reports presence, never values."""
    out: list[Finding] = []
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        return out
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=False, data_only=True)
    props = wb.properties
    for field in ("creator", "lastModifiedBy"):
        if getattr(props, field, None):
            out.append(Finding(path.name, "(workbook)", f"properties.{field}", "REVIEW",
                               "metadata: author field populated (may name a person; value not shown)", None))
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            out.append(Finding(path.name, ws.title, "(sheet)", "REVIEW", f"structure: sheet is {ws.sheet_state}", None))
        hidden_cols = [k for k, d in ws.column_dimensions.items() if d.hidden]
        if hidden_cols:
            out.append(Finding(path.name, ws.title, ",".join(hidden_cols), "REVIEW", "structure: hidden columns", None))
        comments = sum(1 for row in ws.iter_rows() for c in row if c.comment)
        if comments:
            out.append(Finding(path.name, ws.title, "(cells)", "REVIEW", "structure: cell comments present", comments))
    return out


def scan_frame(fname: str, sheet: str, df: pd.DataFrame) -> list[Finding]:
    out: list[Finding] = []
    for col in df.columns:
        header = str(col)
        for sev, label, rx in HEADER_RULES:
            if rx.search(header):
                out.append(Finding(fname, sheet, header, sev, label, None))
        values = df[col].dropna().astype(str).str.strip()
        values = values[values != ""]
        if values.empty:
            continue
        # Cells that are a single number (e.g., 0.3243243243 percentages in tabulated
        # reports) cannot be phone numbers or IPs; excluding them stops false positives.
        non_numeric = values[~values.str.fullmatch(PURE_NUMBER)]
        for sev, label, rx in VALUE_RULES:
            n = sum(1 for v in non_numeric if rx.search(v))
            if n:
                out.append(Finding(fname, sheet, header, sev, label, n))
        mean_len = values.str.len().mean()
        unique_share = values.nunique() / len(values)
        if mean_len >= FREE_TEXT_MIN_MEAN_LEN and unique_share >= FREE_TEXT_MIN_UNIQUE:
            out.append(Finding(fname, sheet, header, "REVIEW",
                               "free text: may contain names or employers; manual review", len(values)))
    return out


def render(files: list[Path], shapes: list[tuple[str, str, int, int]], findings: list[Finding]) -> str:
    high = sum(f.severity == "HIGH" for f in findings)
    review = sum(f.severity == "REVIEW" for f in findings)
    verdict = "CLEAR" if not findings else "FINDINGS — stop (S8); Scott decides case by case"
    lines = [
        "# Identity Scan Report",
        "",
        "Generated by `scripts/p0_scan_identifiers.py`. Contains no cell values.",
        "",
        f"**Result:** {verdict}  ",
        f"**Findings:** {high} HIGH · {review} REVIEW",
        "",
        "## Files scanned",
        "",
        "| File | Sheet | Rows | Columns |",
        "|------|-------|------|---------|",
    ]
    lines += [f"| {f} | {s} | {r} | {c} |" for f, s, r, c in shapes]
    lines += ["", "## Findings", ""]
    if not findings:
        lines.append("None.")
    else:
        lines += ["| Severity | File | Sheet | Column | Check | Cells matched |",
                  "|----------|------|-------|--------|-------|---------------|"]
        order = {"HIGH": 0, "REVIEW": 1}
        for f in sorted(findings, key=lambda x: (order[x.severity], x.file, x.sheet, x.column)):
            cnt = "—" if f.count is None else str(f.count)
            col = f.column.replace("|", "\\|")
            lines.append(f"| {f.severity} | {f.file} | {f.sheet} | {col} | {f.check} | {cnt} |")
    lines += ["", "## Limits", "",
              "- Person names inside free text cannot be detected reliably; free-text columns are always flagged REVIEW.",
              "- Pattern checks can produce false positives (e.g., long numeric codes matching the phone pattern).",
              "- A CLEAR result is a screen, not a guarantee.", ""]
    return "\n".join(lines)


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("files", nargs="+", type=Path)
    ap.add_argument("--out", type=Path, help="write report here (default: stdout)")
    args = ap.parse_args(argv)

    findings: list[Finding] = []
    shapes: list[tuple[str, str, int, int]] = []
    for path in args.files:
        try:
            sheets = read_sheets(path)
            findings += workbook_checks(path)
        except Exception as exc:  # noqa: BLE001 — report and fail closed
            print(f"ERROR reading {path}: {exc}", file=sys.stderr)
            return 2
        for sheet, df in sheets.items():
            shapes.append((path.name, sheet, len(df), len(df.columns)))
            findings += scan_frame(path.name, sheet, df)

    report = render(args.files, shapes, findings)
    if args.out:
        args.out.write_text(report, encoding="utf-8")
    else:
        sys.stdout.write(report)
    return 3 if findings else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
